import os
import re
import json
import docx

def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.strip()
    # Normalize common whitespaces
    text = re.sub(r'[\u00a0\u3000\t]+', ' ', text)
    # Remove trailing page numbers if any
    text = re.sub(r'\s+\d{1,3}$', '', text)
    return text

def extract_questions_from_docx(docx_path: str, output_json_path: str):
    print(f"Loading document: {docx_path}...")
    doc = docx.Document(docx_path)
    print(f"Total paragraphs in document: {len(doc.paragraphs)}")

    current_part = ""
    current_type = ""
    current_chapter = ""
    current_q = None

    questions = []
    subjects_tree = {}

    # Skip TOC (first ~29 paragraphs)
    for i, p in enumerate(doc.paragraphs):
        if i < 29:
            continue
        text = clean_text(p.text)
        if not text:
            continue

        # Detect Part (Subject)
        if any(text.startswith(x) for x in ['第一部分', '第二部分', '第三部分', '第四部分', '第五部分']):
            current_part = text
            if current_part not in subjects_tree:
                subjects_tree[current_part] = {}
            continue

        # Detect Section Type
        if '一、单项选择题' in text:
            current_type = '单项选择题'
            continue
        elif '二、多项选择题' in text:
            current_type = '多项选择题'
            continue
        elif '三、材料分析题' in text:
            current_type = '材料分析题'
            continue

        # Skip material analysis questions as aligned
        if current_type == '材料分析题':
            continue

        # Detect Chapter
        if text.startswith('§ '):
            current_chapter = text[2:].strip()
            if current_part:
                if current_chapter not in subjects_tree[current_part]:
                    subjects_tree[current_part][current_chapter] = {
                        '单项选择题': 0,
                        '多项选择题': 0
                    }
            continue

        # Detect Question Number and Stem
        m_q = re.match(r'^(\d+)\.\s*(.+)', text)
        if m_q and current_type in ('单项选择题', '多项选择题'):
            if current_q:
                questions.append(current_q)
            q_num = int(m_q.group(1))
            stem_content = m_q.group(2).strip()

            q_id = f"Q_{len(questions) + 1:04d}"
            current_q = {
                'id': q_id,
                'part': current_part,
                'chapter': current_chapter,
                'type': current_type,
                'num': q_num,
                'stem': stem_content,
                'sub_items': [],
                'options': []
            }
            if current_part and current_chapter and current_part in subjects_tree and current_chapter in subjects_tree[current_part]:
                subjects_tree[current_part][current_chapter][current_type] += 1
            continue

        # Detect Option
        m_opt = re.match(r'^([A-D])\.\s*(.+)', text)
        if m_opt and current_q:
            current_q['options'].append({
                'key': m_opt.group(1),
                'text': m_opt.group(2).strip()
            })
            continue

        # Detect sub-items like ①②③④⑤
        if re.match(r'^[①②③④⑤⑥⑦⑧⑨⑩]', text) and current_q and not current_q['options']:
            current_q['sub_items'].append(text)
            continue

        # Continuation of previous text
        if current_q:
            if current_q['options']:
                current_q['options'][-1]['text'] += ' ' + text
            elif current_q['sub_items']:
                current_q['sub_items'][-1] += ' ' + text
            else:
                current_q['stem'] += ' ' + text

    if current_q:
        questions.append(current_q)

    # Post-check: ensure option formatting and cleaning
    for q in questions:
        for opt in q['options']:
            opt['text'] = clean_text(opt['text'])
        q['stem'] = clean_text(q['stem'])

    print(f"\nSuccessfully extracted {len(questions)} objective questions!")

    single_count = sum(1 for q in questions if q['type'] == '单项选择题')
    multi_count = sum(1 for q in questions if q['type'] == '多项选择题')
    print(f"- 单项选择题: {single_count} 道")
    print(f"- 多项选择题: {multi_count} 道")

    result = {
        'total': len(questions),
        'single_count': single_count,
        'multi_count': multi_count,
        'subjects_tree': subjects_tree,
        'questions': questions
    }

    os.makedirs(os.path.dirname(output_json_path), exist_ok=True)
    with open(output_json_path, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"Saved database to {output_json_path} (Size: {os.path.getsize(output_json_path) / 1024:.1f} KB)")
    return result

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    parent_dir = os.path.dirname(base_dir)
    docx_file = os.path.join(base_dir, "data", "2027考研政治1000题（试题分册）.docx")
    if not os.path.exists(docx_file):
        docx_file = os.path.join(parent_dir, "2027考研政治1000题（试题分册）.docx")
    target_json = os.path.join(base_dir, "data", "questions.json")
    extract_questions_from_docx(docx_file, target_json)
